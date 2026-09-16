import json
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, RGBColor
from lxml import etree

from app.converter import docx_to_markdown, parse_markdown_blocks, render_markdown_to_docx


def canonical_xml(element) -> str:
    """Serialize OOXML ignoring volatile attributes (rsid, paraId, ...)."""
    xml = etree.tostring(element).decode()
    xml = re.sub(r'xmlns:\w+="[^"]*"\s*', "", xml)
    xml = re.sub(r'\s+(?:w|w14):(?:paraId|textId|rsid\w*)="[^"]*"', "", xml)
    return xml


def make_template(path: Path) -> None:
    """Template exercising the formatting surfaces involved in style fidelity."""
    document = Document()
    mixed = document.add_paragraph()
    bold_run = mixed.add_run("Bold intro ")
    bold_run.bold = True
    bold_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    plain_run = mixed.add_run("and regular text")
    plain_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    spaced = document.add_paragraph()
    spaced.paragraph_format.space_after = Inches(0.3)
    spaced.add_run("Second paragraph")
    blank = document.add_paragraph()
    blank.add_run().add_break(WD_BREAK.PAGE)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Filled"
    # second cell stays empty on purpose
    document.save(path)


def round_trip(source: Path, markdown: Path, manifest: Path, output: Path, edited: str | None = None):
    docx_to_markdown(source, markdown, manifest)
    text = markdown.read_text(encoding="utf-8") if edited is None else edited
    return render_markdown_to_docx(source, text, manifest, output), Document(output)


class StyleFidelityTests(unittest.TestCase):
    def test_round_trip_preserves_per_run_formatting(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            _, exported = round_trip(source, markdown, manifest, output)
            runs = exported.paragraphs[0].runs
            self.assertEqual(len(runs), 2)
            self.assertTrue(runs[0].bold)
            self.assertFalse(runs[1].bold)
            self.assertEqual(runs[0].font.color.rgb, RGBColor(0xC0, 0x00, 0x00))
            self.assertEqual(runs[1].font.color.rgb, RGBColor(0x66, 0x66, 0x66))

    def test_unedited_body_is_xml_identical_to_template(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            _, exported = round_trip(source, markdown, manifest, output)
            template_children = [c for c in Document(str(source)).element.body]
            export_children = [c for c in exported.element.body]
            self.assertEqual(len(template_children), len(export_children))
            for expected, actual in zip(template_children, export_children):
                self.assertEqual(canonical_xml(expected), canonical_xml(actual))

    def test_page_break_only_paragraph_keeps_inline_break(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            _, exported = round_trip(source, markdown, manifest, output)
            target = None
            for paragraph in exported.paragraphs:
                if paragraph._p.xpath('.//w:br[@w:type="page"]'):
                    target = paragraph
                    break
            self.assertIsNotNone(target, "inline page break lost in export")
            self.assertFalse(bool(target._p.xpath(".//w:pageBreakBefore")))
            self.assertTrue(bool(target._p.xpath(".//w:r")))

    def test_no_empty_runs_are_added(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            _, exported = round_trip(source, markdown, manifest, output)
            for paragraph in exported.paragraphs:
                if paragraph.text:
                    continue
                if paragraph._p.xpath('.//w:br[@w:type="page"]'):
                    continue
                self.assertEqual(
                    paragraph._p.xpath(".//w:r"),
                    [],
                    f"empty run added to paragraph {paragraph.text!r}",
                )
            cell = exported.tables[0].cell(0, 1)
            for paragraph in cell.paragraphs:
                self.assertEqual(paragraph._p.xpath(".//w:r"), [])

    def test_edited_paragraph_does_not_inherit_first_run_bold(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            docx_to_markdown(source, markdown, manifest)
            text = markdown.read_text(encoding="utf-8").replace(
                "**Bold intro **and regular text", "Fully rewritten plain text"
            )
            self.assertIn("Fully rewritten plain text", text)
            _, exported = round_trip(source, markdown, manifest, output, edited=text)
            rewritten = exported.paragraphs[0]
            self.assertEqual(rewritten.text, "Fully rewritten plain text")
            for run in rewritten.runs:
                self.assertFalse(run.bold, f"bold bled onto run {run.text!r}")

    def test_edited_paragraph_does_not_write_explicit_page_break_before(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            docx_to_markdown(source, markdown, manifest)
            text = markdown.read_text(encoding="utf-8").replace("Second paragraph", "Rewritten paragraph")
            result, exported = round_trip(source, markdown, manifest, output, edited=text)
            self.assertFalse(result["warnings"])
            for paragraph in exported.paragraphs:
                if "Rewritten paragraph" in paragraph.text:
                    page_breaks = paragraph._p.xpath(".//w:pageBreakBefore")
                    self.assertEqual(page_breaks, [], "explicit pageBreakBefore written on normal paragraph")
                    return
            self.fail("edited paragraph not found in export")

    def test_unedited_table_is_xml_identical_to_template(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Filled"
            # Word often leaves empty runs carrying run properties behind
            leftover = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            properties.append(OxmlElement("w:b"))
            leftover.append(properties)
            table.cell(0, 0).paragraphs[0]._p.append(leftover)
            document.save(source)
            _, exported = round_trip(source, markdown, manifest, output)
            template_table = Document(str(source)).tables[0]._tbl
            export_table = exported.tables[0]._tbl
            self.assertEqual(canonical_xml(template_table), canonical_xml(export_table))

    def test_standalone_page_break_block_uses_page_break_run(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, markdown, manifest, output = (root / name for name in ("t.docx", "w.md", "m.json", "e.docx"))
            make_template(source)
            docx_to_markdown(source, markdown, manifest)
            text = markdown.read_text(encoding="utf-8") + "\n<!-- dc:page-break -->\nAfter the break\n"
            _, exported = round_trip(source, markdown, manifest, output, edited=text)
            page_break_paragraphs = [p for p in exported.paragraphs if p._p.xpath('.//w:br[@w:type="page"]')]
            self.assertGreaterEqual(len(page_break_paragraphs), 2)
            standalone = page_break_paragraphs[-1]
            self.assertFalse(bool(standalone._p.xpath(".//w:pageBreakBefore")))
            self.assertEqual(standalone.text, "")
            self.assertTrue(parse_markdown_blocks("x\n\n<!-- dc:page-break -->\n"))


class ConverterTests(unittest.TestCase):
    def test_template_round_trip_keeps_template_parts_and_updates_content(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "template.docx"
            markdown = root / "working.md"
            manifest = root / "manifest.json"
            output = root / "export.docx"

            document = Document()
            section = document.sections[0]
            section.left_margin = Inches(1.1)
            section.right_margin = Inches(0.8)
            section.header.paragraphs[0].text = "ACME SECURITY · CONFIDENTIAL"
            section.footer.paragraphs[0].text = "Pre-engagement"
            document.add_heading("Web API Pre-Engagement", level=1)
            document.add_paragraph("Client: Example Corp", style="Subtitle")
            document.add_paragraph("Scope details go here.", style="Body Text")
            document.add_paragraph("Auth review", style="List Bullet")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Asset"
            table.cell(0, 1).text = "Environment"
            table.cell(1, 0).text = "api.example.test"
            table.cell(1, 1).text = "Staging"
            document.add_section(WD_SECTION.NEW_PAGE)
            document.add_heading("Second section", level=1)
            document.save(source)

            result = docx_to_markdown(source, markdown, manifest)
            self.assertGreaterEqual(len(result["editable_blocks"]), 6)
            source_markdown = markdown.read_text(encoding="utf-8")
            self.assertIn("dc:block", source_markdown)
            self.assertIn("Web API Pre-Engagement", source_markdown)
            self.assertIn("| Asset | Environment |", source_markdown)

            edited = source_markdown.replace("Example Corp", "Northwind Labs")
            edited = edited.replace("Scope details go here.", "The API scope includes the partner gateway and its authentication flows.")
            render_result = render_markdown_to_docx(source, edited, manifest, output)
            self.assertFalse(render_result["warnings"])

            exported = Document(output)
            text = "\n".join(p.text for p in exported.paragraphs)
            self.assertIn("Northwind Labs", text)
            self.assertIn("partner gateway", text)
            self.assertIn("Second section", text)
            self.assertEqual(exported.sections[0].header.paragraphs[0].text, "ACME SECURITY · CONFIDENTIAL")
            self.assertEqual(exported.sections[0].footer.paragraphs[0].text, "Pre-engagement")
            self.assertAlmostEqual(exported.sections[0].left_margin.inches, 1.1, places=2)
            self.assertEqual(exported.tables[0].cell(1, 0).text, "api.example.test")
            self.assertEqual(json.loads(manifest.read_text(encoding="utf-8"))["version"], 1)


if __name__ == "__main__":
    unittest.main()

